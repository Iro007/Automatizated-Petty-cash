"""Reproducible Streamlit AppTest checks using synthetic expenses and BCV HTML."""
from datetime import date
from email.message import Message
from pathlib import Path
import re
from unittest.mock import patch
import json
import math
import sys

import openpyxl
import pandas as pd
import streamlit as st
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from excel_builder import build_caja_chica
from i18n import tr

OUT = ROOT / "evidence/demo_20261003"
OUT.mkdir(parents=True, exist_ok=True)
BCV_HTML = b"""<div id='dolar'><strong class='strong-tb'>871,36890000</strong></div>
<span class='date-display-single' content='2026-10-05T00:00:00-04:00'>Lunes, 05 Octubre 2026</span>"""


class FakeResponse:
    status = 200

    def __init__(self, body=BCV_HTML):
        self.body = body
        self.headers = Message()
        self.headers["Content-Type"] = "text/html; charset=utf-8"

    def read(self):
        return self.body

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False


checks = []


def record(name, result, details=None):
    checks.append({"name": name, "passed": bool(result), "details": details})


def app():
    return AppTest.from_file(str(ROOT / "app.py"), default_timeout=30)


def text_input(at, label):
    return next(item for item in at.text_input if item.label == label)


def action(at, key):
    label = tr(at.session_state["language"], key)
    return next(item for item in at.button if item.label == label)


def method_selector(at):
    return at.radio(key=f'upload_method_{at.session_state["language"]}')


st.cache_data.clear()
bcv_mock = patch("bcv_rate.urlopen", side_effect=lambda *_args, **_kwargs: FakeResponse())
urlopen_mock = bcv_mock.start()

at = app()
at.session_state["language"] = "es"
at.session_state["language_last"] = "es"
at.run()
record("app_initial_render", len(at.exception) == 0, [str(x.value) for x in at.exception])
record(
    "bcv_loaded_on_first_session_run",
    at.session_state["bcv_reference_rate"] == "871.36890000"
    and at.session_state["bcv_value_date"] == "2026-10-05",
    {"rate": at.session_state["bcv_reference_rate"], "value_date": at.session_state["bcv_value_date"]},
)
record("initial_screen_has_no_premature_validation_errors", len(at.error) == 0, [x.value for x in at.error])
record(
    "empty_expenses_disable_generation",
    action(at, "generate_excel").disabled and at.session_state["last_file"] is None,
    {"disabled": action(at, "generate_excel").disabled},
)
record(
    "receipt_method_is_default_and_hides_manual_fields",
    method_selector(at).value == "receipts"
    and not any(item.key == "manual_supplier" for item in at.text_input),
    {"method": method_selector(at).value, "text_input_keys": [item.key for item in at.text_input]},
)
issue_date_field = text_input(at, "Fecha de emisión (dd/mm/aaaa)")
record("issue_date_defaults_to_venezuelan_format", re.fullmatch(r"\d{2}/\d{2}/\d{4}", issue_date_field.value) is not None, issue_date_field.value)

at.run()
record("bcv_fetch_is_cached_during_session", urlopen_mock.call_count == 1, urlopen_mock.call_count)
action(at, "rate_refresh").click().run()
record("explicit_bcv_refresh_bypasses_cache", urlopen_mock.call_count == 2, urlopen_mock.call_count)

receipt_method_states = []
at.radio(key="language").set_value("en").run()
receipt_method_states.append(method_selector(at).value == "receipts" and at.session_state["input_method"] == "receipts")
at.toggle(key="dark_mode").set_value(True).run()
receipt_method_states.append(method_selector(at).value == "receipts" and not any(item.key == "manual_supplier" for item in at.text_input))
at.radio(key="language").set_value("es").run()
receipt_method_states.append(method_selector(at).value == "receipts" and at.session_state["input_method"] == "receipts")
at.toggle(key="dark_mode").set_value(False).run()
record("receipt_method_survives_language_and_theme_switch", all(receipt_method_states), receipt_method_states)

method_selector(at).set_value("manual").run()
at.radio(key="language").set_value("en").run()
record(
    "manual_language_switch_translates_ui",
    any(x.label == "Vendor / Merchant*" for x in at.text_input)
    and any(x.label == "Language" for x in at.radio),
    [x.label for x in at.text_input],
)
next(x for x in at.toggle if x.label == "Dark mode").set_value(True).run()
css = "\n".join(str(item.value) for item in at.markdown)
record(
    "dark_theme_toggle_applies_dark_palette",
    at.session_state["dark_mode"] and "--app-bg:#0F172A" in css,
    {"dark_mode": at.session_state["dark_mode"], "dark_css_present": "--app-bg:#0F172A" in css},
)
next(x for x in at.toggle if x.label == "Dark mode").set_value(False).run()
at.radio(key="language").set_value("es").run()

action(at, "manual_add_button").click().run()
record("empty_manual_entry_rejected", any("obligatorios" in x.value for x in at.error), [x.value for x in at.error])
text_input(at, "Proveedor / Comercio*").set_value("Proveedor ficticio")
text_input(at, "Fecha (dd/mm/aaaa)").set_value("31/02/2026")
text_input(at, "Monto (Bs)*").set_value("10,00")
action(at, "manual_add_button").click().run()
record(
    "invalid_manual_date_rejected",
    len(at.session_state["gastos"]) == 0 and any("Fecha inválida" in x.value for x in at.error),
    [x.value for x in at.error],
)
record(
    "invalid_manual_submission_keeps_entered_values",
    at.text_input(key="manual_supplier").value == "Proveedor ficticio"
    and at.text_input(key="manual_expense_date").value == "31/02/2026"
    and at.text_input(key="manual_amount").value == "10,00",
    {item.key: item.value for item in at.text_input if item.key.startswith("manual_")},
)
manual_method_states = []
at.radio(key="language").set_value("en").run()
manual_method_states.append(method_selector(at).value == "manual" and at.session_state["input_method"] == "manual")
at.toggle(key="dark_mode").set_value(True).run()
manual_method_states.append(method_selector(at).value == "manual" and at.text_input(key="manual_supplier").value == "Proveedor ficticio")
at.radio(key="language").set_value("es").run()
manual_method_states.append(method_selector(at).value == "manual" and at.session_state["input_method"] == "manual")
at.toggle(key="dark_mode").set_value(False).run()
record(
    "manual_method_and_draft_survive_language_and_theme_switch",
    all(manual_method_states)
    and at.text_input(key="manual_expense_date").value == "31/02/2026"
    and at.text_input(key="manual_amount").value == "10,00",
    {"method_states": manual_method_states, "draft": {item.key: item.value for item in at.text_input if item.key.startswith("manual_")}},
)
method_selector(at).set_value("receipts").run()
record(
    "receipts_method_hides_manual_draft",
    not any(item.key == "manual_supplier" for item in at.text_input),
    [item.key for item in at.text_input],
)
method_selector(at).set_value("manual").run()
record(
    "manual_draft_survives_method_switch",
    at.text_input(key="manual_supplier").value == "Proveedor ficticio"
    and at.text_input(key="manual_expense_date").value == "31/02/2026"
    and at.text_input(key="manual_amount").value == "10,00",
    {item.key: item.value for item in at.text_input if item.key.startswith("manual_")},
)
text_input(at, "Fecha (dd/mm/aaaa)").set_value("03/10/2026")
text_input(at, "Monto (Bs)*").set_value("10.00")
action(at, "manual_add_button").click().run()
record(
    "manual_period_decimal_accepted",
    len(at.session_state["gastos"]) == 1
    and math.isclose(float(at.session_state["gastos"].iloc[0]["monto_bs"]), 10.0),
    at.session_state["gastos"].to_dict(orient="records"),
)
record(
    "successful_manual_submission_clears_draft",
    at.text_input(key="manual_supplier").value == ""
    and at.text_input(key="manual_amount").value == "",
    {"supplier": at.text_input(key="manual_supplier").value, "amount": at.text_input(key="manual_amount").value},
)
text_input(at, "Proveedor / Comercio*").set_value("Segundo gasto")
text_input(at, "Monto (Bs)*").set_value("10,00")
action(at, "manual_add_button").click().run()
record(
    "manual_comma_decimal_accepted",
    len(at.session_state["gastos"]) == 2
    and math.isclose(float(at.session_state["gastos"].iloc[1]["monto_bs"]), 10.0),
    at.session_state["gastos"].to_dict(orient="records"),
)

rows = [
    {"proveedor": "Papelería de ejemplo", "fecha": "03/10/2026", "factura": "100001", "monto_bs": 1250.50, "descripcion": "Materiales ficticios"},
    {"proveedor": "Transporte de ejemplo", "fecha": "03/10/2026", "factura": "100002", "monto_bs": 749.50, "descripcion": "Monto revisado contra comprobante sintético"},
    {"proveedor": "Servicio de ejemplo", "fecha": "03/10/2026", "factura": "100003", "monto_bs": 500.0, "descripcion": "Gasto ficticio agregado manualmente"},
]
at.session_state["gastos"] = pd.DataFrame(rows)
at.session_state["editor_version"] += 1
for item in at.text_input:
    values = {
        "Nombre": "DEMOSTRACIÓN FICTICIA",
        "RIF": "SIN VALIDEZ FISCAL",
        "Dirección": "Datos sintéticos",
        "Teléfonos": "",
        "Responsable": "Responsable de ejemplo",
        "Cédula": "DEMO-000001",
        "Monto otorgado ($)": "100,00",
        "Tasa aplicada (Bs/USD)": "50,00000000",
        "Fecha de emisión (dd/mm/aaaa)": "03/10/2026",
        "N° Reporte": "DEMO-PRUEBAS-20261003",
    }
    if item.label in values:
        item.set_value(values[item.label])
at.run()
text_input(at, "Fecha de emisión (dd/mm/aaaa)").set_value("31/02/2026").run()
action(at, "generate_excel").click().run()
record(
    "invalid_issue_date_blocks_generation",
    any("Fecha de emisión inválida" in item.value for item in at.warning)
    and at.session_state["last_file"] is None,
    [item.value for item in at.warning],
)
text_input(at, "Fecha de emisión (dd/mm/aaaa)").set_value("03/10/2026").run()
at.session_state["gastos"] = pd.DataFrame([{**rows[0], "fecha": "31/02/2026"}])
at.session_state["editor_version"] += 1
at.run()
action(at, "generate_excel").click().run()
record(
    "invalid_edited_date_blocks_generation",
    any("Gasto 1: Fecha inválida" in x.value for x in at.warning)
    and at.session_state["last_file"] is None,
    [x.value for x in at.warning],
)
at.session_state["gastos"] = pd.DataFrame(rows)
at.session_state["editor_version"] += 1
at.run()
record(
    "corrected_date_reenables_generation",
    not any("Fecha inválida" in x.value for x in at.error)
    and not action(at, "generate_excel").disabled,
    [x.value for x in at.error],
)
action(at, "generate_excel").click().run()
record("app_generates_workbook", len(at.exception) == 0 and at.session_state["last_file"] is not None, [str(x.value) for x in at.exception])
_name, data, source_path = at.session_state["last_file"]
example = OUT / "caja_chica_ejemplo.xlsx"
example.write_bytes(data)
wb = openpyxl.load_workbook(example, data_only=False)
ws = wb.active
record(
    "actual_workbook_values",
    [ws[f"F{r}"].value for r in range(11, 14)] == [1250.50, 749.50, 500.0]
    and ws["F8"].value == 50
    and ws["C8"].value == 100,
    {"amounts": [ws[f"F{r}"].value for r in range(11, 14)], "rate": ws["F8"].value, "fund": ws["C8"].value},
)
formulas = {coord: ws[coord].value for coord in ["G11", "G12", "G13", "F14", "G14", "G15"]}
record(
    "formula_references",
    formulas == {
        "G11": "=IF($F$8=0,0,F11/$F$8)",
        "G12": "=IF($F$8=0,0,F12/$F$8)",
        "G13": "=IF($F$8=0,0,F13/$F$8)",
        "F14": "=SUM(F11:F13)",
        "G14": "=SUM(G11:G13)",
        "G15": "=$C$8-G14",
    },
    formulas,
)
record(
    "manual_adjustment_metadata_saved",
    "ajuste manual" in ws["B9"].value.lower()
    and "871,36890000" in ws["B9"].value
    and "05/10/2026" in ws["B9"].value,
    ws["B9"].value,
)
record("rate_remains_numeric_with_eight_decimals", isinstance(ws["F8"].value, (float, int)) and "0.00000000" in ws["F8"].number_format, ws["F8"].number_format)
record("independent_totals", math.isclose(sum(r["monto_bs"] for r in rows), 2500) and math.isclose(sum(r["monto_bs"] for r in rows) / 50, 50), {"bs": 2500, "usd": 50, "balance": 50})

text_input(at, "Tasa aplicada (Bs/USD)").set_value("55,00000000").run()
record(
    "editing_rate_invalidates_existing_download",
    at.session_state["last_file"] is None,
    {"rate": at.session_state["rate_input"], "last_file_cleared": at.session_state["last_file"] is None},
)
text_input(at, "Tasa aplicada (Bs/USD)").set_value("50,00000000").run()
action(at, "generate_excel").click().run()
text_input(at, "N° Reporte").set_value("DEMO-PRUEBAS-EDITADO").run()
record(
    "editing_report_invalidates_existing_download",
    at.session_state["last_file"] is None,
    {"report": at.session_state["n_reporte"], "last_file_cleared": at.session_state["last_file"] is None},
)
text_input(at, "N° Reporte").set_value("DEMO-PRUEBAS-20261003").run()

at.session_state["language"] = "en"
at.session_state["language_last"] = "es"
at.session_state["editor_version"] += 1
at.run()
for item in at.text_input:
    if item.label == "Amount granted (USD)": item.set_value("100.00")
    if item.label == "Applied rate (Bs/USD)": item.set_value("50.00000000")
at.run()
record("english_amount_rate_and_date_inputs_render", any(item.label == "Amount granted (USD)" for item in at.text_input) and any(item.label == "Applied rate (Bs/USD)" for item in at.text_input) and any(item.label == "Issue date (DD/MM/YYYY)" for item in at.text_input), [item.label for item in at.text_input])
action(at, "generate_excel").click().run()
_, english_data, _ = at.session_state["last_file"]
english_example = OUT / "petty_cash_example.xlsx"
english_example.write_bytes(english_data)
english_export = openpyxl.load_workbook(english_example, data_only=False).active
record(
    "english_app_export_translates_workbook_and_keeps_rate_provenance",
    english_export["B5"].value == "PETTY CASH EXPENSE REPORT"
    and english_export["C10"].value == "VENDOR / MERCHANT"
    and english_export["F8"].value == 50
    and "Original BCV reference: 871.36890000" in english_export["B9"].value,
    {"title": english_export["B5"].value, "vendor_header": english_export["C10"].value, "rate": english_export["F8"].value, "provenance": english_export["B9"].value},
)
english_book = build_caja_chica(rows, idioma="en", fecha_emision=date(2026, 10, 3), tasa_bcv=50, monto_otorgado_usd=100)
record("english_workbook_header", english_book.active["B10"].value == "ITEM" and english_book.active["C10"].value == "VENDOR / MERCHANT", [english_book.active["B10"].value, english_book.active["C10"].value])

urlopen_mock.side_effect = TimeoutError("synthetic offline test")
st.cache_data.clear()
offline_at = app()
offline_at.session_state["language"] = "es"
offline_at.session_state["language_last"] = "es"
offline_at.run()
record(
    "bcv_failure_shows_warning_and_manual_fallback",
    offline_at.session_state["bcv_last_error"]
    and offline_at.session_state["bcv_reference_rate"] is None
    and any("No se pudo consultar el BCV" in item.value for item in offline_at.warning),
    {"bcv_last_error": offline_at.session_state["bcv_last_error"], "warnings": [item.value for item in offline_at.warning]},
)
text_input(offline_at, "Tasa aplicada (Bs/USD)").set_value("123,45678901")
offline_at.run()
record(
    "manual_rate_is_not_labeled_as_current_bcv",
    offline_at.session_state["rate_source"] == "manual"
    and offline_at.session_state["bcv_reference_rate"] is None,
    {"rate_source": offline_at.session_state["rate_source"], "rate": offline_at.session_state["rate_input"]},
)

urlopen_mock.side_effect = lambda *_args, **_kwargs: FakeResponse()
st.cache_data.clear()
stale_at = app()
stale_at.session_state["language"] = "es"
stale_at.session_state["language_last"] = "es"
stale_at.run()
urlopen_mock.side_effect = TimeoutError("synthetic refresh failure")
action(stale_at, "rate_refresh").click().run()
record(
    "failed_refresh_marks_retained_bcv_value_stale",
    stale_at.session_state["bcv_last_error"]
    and stale_at.session_state["rate_source"] == "bcv_stale"
    and any("Último dato BCV recibido" in item.value for item in stale_at.caption),
    {"source": stale_at.session_state["rate_source"], "captions": [item.value for item in stale_at.caption]},
)
urlopen_mock.side_effect = lambda *_args, **_kwargs: FakeResponse()

for count in [0, 1, 30, 1000]:
    workbook = build_caja_chica(rows[:1] * count, tasa_bcv=50, monto_otorgado_usd=100)
    sheet = workbook.active
    end = 10 + max(count, 1)
    total = end + 1
    record(f"row_boundary_{count}", sheet[f"F{total}"].value == f"=SUM(F11:F{end})" and sheet[f"G{end}"].value == f"=IF($F$8=0,0,F{end}/$F$8)", {"actual_rows": count, "total_row": total, "sum_formula": sheet[f"F{total}"].value})

for item in at.text_input:
    if item.label == "Applied rate (Bs/USD)": item.set_value("0")
at.run()
action(at, "generate_excel").click().run()
record(
    "zero_rate_blocks_generation",
    any("Bs/USD rate" in item.value for item in at.warning)
    and at.session_state["last_file"] is None,
    [item.value for item in at.warning],
)

result = {
    "checks": checks,
    "all_checks_passed": all(item["passed"] for item in checks),
    "example": example.name,
    "source_generated_path": str(source_path),
    "bcv_fetch_count_including_explicit_refresh": urlopen_mock.call_count,
    "native_excel_recalculation": "not_verified",
    "rate": "Synthetic BCV HTML fixture, not a current exchange rate",
}
(OUT / "test_results.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
bcv_mock.stop()
print(json.dumps(result, ensure_ascii=False, indent=2))
if not result["all_checks_passed"]:
    raise SystemExit(1)
