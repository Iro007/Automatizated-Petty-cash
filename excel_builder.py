"""Generador propio de planilla Caja Chica (.xlsx).

No depende de ningún archivo base externo: construye el workbook
desde cero con openpyxl, con estilos, fórmulas, firmas y
configuración de impresión incluidos.

Uso:
    from excel_builder import build_caja_chica
    wb = build_caja_chica(df, empresa={...}, responsable={...}, ...)
"""
from __future__ import annotations

from copy import copy
from datetime import date, datetime
import re
from typing import Iterable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from i18n import excel_number_format, format_number, normalize_language, tr

# ---------------- Paleta propia ----------------
NAVY = "1F3864"
BLUE = "2E75B6"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
DARK_TEXT = "1F1F1F"
GRAY_TEXT = "595959"

TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=WHITE)
SUBTITLE_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=10, color=DARK_TEXT)
SMALL_FONT = Font(name="Calibri", size=9, color=GRAY_TEXT)
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color=DARK_TEXT)
TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)

FILL_NAVY = PatternFill("solid", fgColor=NAVY)
FILL_BLUE = PatternFill("solid", fgColor=BLUE)
FILL_LIGHT_BLUE = PatternFill("solid", fgColor=LIGHT_BLUE)
FILL_LIGHT_GRAY = PatternFill("solid", fgColor=LIGHT_GRAY)
FILL_WHITE = PatternFill("solid", fgColor=WHITE)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER_NOWRAP = Alignment(horizontal="center", vertical="center", wrap_text=False)

THIN = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_DATE = "DD/MM/YYYY"


def validar_fecha_gasto(valor: date | str, idioma: str = "es") -> str:
    """Return a real calendar date in dd/mm/yyyy, or reject it explicitly."""
    mensaje = tr(idioma, "invalid_date")
    if isinstance(valor, date):
        try:
            texto = f"{valor.day:02d}/{valor.month:02d}/{valor.year:04d}"
        except (ValueError, TypeError):
            raise ValueError(mensaje) from None
    elif isinstance(valor, str):
        texto = valor.strip()
    else:
        raise ValueError(mensaje)
    if not re.fullmatch(r"[0-9]{2}/[0-9]{2}/[0-9]{4}", texto):
        raise ValueError(mensaje)
    try:
        datetime.strptime(texto, "%d/%m/%Y")
    except ValueError:
        raise ValueError(mensaje) from None
    return texto


def _escribir_texto_literal(cell, valor):
    """Store user text as an XLSX string, even if it starts with '='."""
    cell.value = str(valor or "")
    cell.data_type = "s"


def _style_range(ws, row: int, col_start: int, col_end: int, font=None, fill=None, alignment=None, border=None, number_format=None):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        if font:
            cell.font = copy(font)
        if fill:
            cell.fill = copy(fill)
        if alignment:
            cell.alignment = copy(alignment)
        if border:
            cell.border = copy(border)
        if number_format:
            cell.number_format = number_format


def build_caja_chica(
    gastos: Iterable[dict] | object,
    empresa: dict | None = None,
    responsable: str = "",
    cedula: str = "",
    fecha_emision: date | str | None = None,
    monto_otorgado_usd: float = 0.0,
    tasa_bcv: float = 0.0,
    n_reporte: str = "",
    titulo: str = "RELACIÓN DE GASTOS — CAJA CHICA",
    idioma: str = "es",
    fecha_valor_bcv: date | str | None = None,
    origen_tasa: str = "manual",
    tasa_bcv_original: float | str | None = None,
    bcv_actualizacion_fallida: bool = False,
) -> openpyxl.Workbook:
    """Construye el workbook de caja chica.

    gastos: lista de dicts o DataFrame con keys/cols:
        proveedor, fecha (dd/mm/yyyy o date), factura, monto_bs (float), descripcion
    """
    lang = normalize_language(idioma)
    titulo_default = "RELACIÓN DE GASTOS — CAJA CHICA"
    if titulo == titulo_default and lang == "en":
        titulo = tr(lang, "excel_default_title")

    empresa = empresa or {}
    nombre = empresa.get("nombre", tr(lang, "default_company_name"))
    rif = empresa.get("rif", tr(lang, "default_rif"))
    direccion = empresa.get("direccion", tr(lang, "default_address"))
    telefonos = empresa.get("telefonos", tr(lang, "default_phones"))

    # Normalizar gastos a lista de dicts
    rows: list[dict] = []
    try:
        import pandas as pd  # type: ignore

        if isinstance(gastos, pd.DataFrame):
            for _, r in gastos.iterrows():
                rows.append(
                    {
                        "proveedor": str(r.get("proveedor", "") or ""),
                        "fecha": r.get("fecha", ""),
                        "factura": str(r.get("factura", "") or ""),
                        "monto_bs": float(r.get("monto_bs", 0) or 0),
                        "descripcion": str(r.get("descripcion", "") or ""),
                    }
                )
        else:
            for g in gastos:  # type: ignore
                rows.append(
                    {
                        "proveedor": str(g.get("proveedor", "") or ""),
                        "fecha": g.get("fecha", ""),
                        "factura": str(g.get("factura", "") or ""),
                        "monto_bs": float(g.get("monto_bs", 0) or 0),
                        "descripcion": str(g.get("descripcion", "") or ""),
                    }
                )
    except Exception:
        rows = list(gastos) if isinstance(gastos, list) else []

    for idx, g in enumerate(rows, start=1):
        try:
            g["fecha"] = validar_fecha_gasto(g.get("fecha", ""), lang)
        except ValueError as exc:
            raise ValueError(tr(lang, "excel_expense_prefix", number=idx, error=exc)) from None
    if fecha_emision is None:
        fecha_emision = date.today()
    else:
        try:
            fecha_emision = datetime.strptime(validar_fecha_gasto(fecha_emision, lang), "%d/%m/%Y").date()
        except ValueError as exc:
            raise ValueError(tr(lang, "excel_issue_date_prefix", error=exc)) from None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tr(lang, "excel_sheet_title")
    fmt_bs = excel_number_format(lang, 2, "Bs")
    fmt_usd = excel_number_format(lang, 2, "USD")
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.center.text = f"&9&I{titulo}"
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)

    # Anchos propios: B..H (7 columnas útiles + A margen)
    widths = {"A": 2, "B": 12, "C": 34, "D": 14, "E": 18, "F": 17, "G": 15, "H": 36}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # --- Encabezado empresa (filas 1-4) ---
    ws.merge_cells("B1:H1")
    _escribir_texto_literal(ws["B1"], nombre.upper())
    ws["B1"].font = Font(name="Calibri", size=17, bold=True, color=NAVY)
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    for r, val in ((2, rif), (3, direccion), (4, telefonos)):
        ws.merge_cells(f"B{r}:H{r}")
        _escribir_texto_literal(ws[f"B{r}"], val)
        ws[f"B{r}"].font = SMALL_FONT
        ws[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 14

    # --- Cinta de título (fila 5) ---
    ws.merge_cells("B5:H5")
    _escribir_texto_literal(ws["B5"], titulo)
    ws["B5"].font = TITLE_FONT
    ws["B5"].alignment = CENTER
    _style_range(ws, 5, 2, 8, fill=FILL_NAVY)
    ws.row_dimensions[5].height = 26

    # --- Bloque de datos (filas 6-8) ---
    ws.merge_cells("B6:H6")
    ws["B6"] = tr(
        lang,
        "excel_report_line",
        report=n_reporte,
        date=fecha_emision.strftime("%d/%m/%Y") if isinstance(fecha_emision, date) else fecha_emision,
    )
    ws["B6"].font = Font(name="Calibri", size=10, bold=True, color=NAVY)
    ws["B6"].alignment = CENTER
    _style_range(ws, 6, 2, 8, fill=FILL_LIGHT_BLUE)
    ws.row_dimensions[6].height = 18

    rate_label_key = {
        "bcv": "excel_rate_bcv",
        "bcv_stale": "excel_rate_bcv_stale",
        "ajuste_manual": "excel_rate_adjusted",
    }.get(origen_tasa, "excel_rate_manual")
    labels = [
        ("B7", tr(lang, "excel_responsible")),
        ("B8", tr(lang, "excel_granted")),
        ("E7", tr(lang, "excel_id")),
        ("E8", tr(lang, rate_label_key)),
    ]
    for coord, txt in labels:
        ws[coord] = txt
        ws[coord].font = Font(name="Calibri", size=9, bold=True, color=GRAY_TEXT)
        ws[coord].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    ws.merge_cells("C7:D7")
    _escribir_texto_literal(ws["C7"], responsable)
    ws["C7"].font = BOLD_FONT
    ws["C7"].alignment = LEFT
    ws["C7"].border = THIN_BORDER

    ws.merge_cells("F7:H7")
    _escribir_texto_literal(ws["F7"], cedula)
    ws["F7"].font = BODY_FONT
    ws["F7"].alignment = LEFT
    ws["F7"].border = THIN_BORDER
    ws["F7"].number_format = "@"

    ws.merge_cells("C8:D8")
    ws["C8"] = float(monto_otorgado_usd or 0)
    ws["C8"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    ws["C8"].alignment = CENTER
    ws["C8"].border = THIN_BORDER
    ws["C8"].number_format = fmt_usd

    ws.merge_cells("F8:H8")
    ws["F8"] = float(tasa_bcv or 0)
    ws["F8"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    ws["F8"].alignment = CENTER
    ws["F8"].border = THIN_BORDER
    ws["F8"].number_format = excel_number_format(lang, 8)
    ws.row_dimensions[7].height = 28
    ws.row_dimensions[8].height = 36

    applied_rate = format_number(float(tasa_bcv or 0), lang, 8)
    if origen_tasa == "bcv_stale" and fecha_valor_bcv:
        fecha_meta = validar_fecha_gasto(fecha_valor_bcv, lang)
        metadata = tr(
            lang,
            "excel_rate_bcv_stale_meta",
            applied=applied_rate,
            date=fecha_meta,
        )
    elif origen_tasa == "bcv" and fecha_valor_bcv:
        fecha_meta = validar_fecha_gasto(fecha_valor_bcv, lang)
        metadata = tr(lang, "excel_rate_bcv_meta", date=fecha_meta)
    elif origen_tasa == "ajuste_manual" and fecha_valor_bcv and tasa_bcv_original is not None:
        fecha_meta = validar_fecha_gasto(fecha_valor_bcv, lang)
        original_rate = format_number(float(tasa_bcv_original), lang, 8)
        metadata = tr(
            lang,
            "excel_rate_adjusted_failed_meta" if bcv_actualizacion_fallida else "excel_rate_adjusted_meta",
            applied=applied_rate,
            rate=original_rate,
            date=fecha_meta,
        )
    else:
        metadata = tr(lang, "excel_rate_manual_meta")
    ws.merge_cells("B9:H9")
    _escribir_texto_literal(ws["B9"], metadata)
    ws["B9"].font = SMALL_FONT
    ws["B9"].alignment = LEFT
    ws["B9"].fill = copy(FILL_LIGHT_GRAY)
    ws.row_dimensions[9].height = 30

    # --- Cabecera de tabla (fila 10) ---
    headers = [
        tr(lang, "excel_item"),
        tr(lang, "excel_supplier"),
        tr(lang, "excel_date"),
        tr(lang, "excel_invoice"),
        tr(lang, "excel_total_bs"),
        tr(lang, "excel_total_usd"),
        tr(lang, "excel_description"),
    ]
    ws.row_dimensions[10].height = 24
    for i, h in enumerate(headers, start=2):
        cell = ws.cell(row=10, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = copy(FILL_BLUE)
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    start_row = 11
    n = max(len(rows), 1)
    end_row = start_row + n - 1

    # --- Cuerpo ---
    for idx in range(n):
        r = start_row + idx
        ws.row_dimensions[r].height = 22
        g = rows[idx] if idx < len(rows) else {"proveedor": "", "fecha": "", "factura": "", "monto_bs": None, "descripcion": ""}
        # Reserve enough height for wrapped descriptions in the generated sheet.
        description_lines = max(1, (len(str(g.get("descripcion", "") or "")) + 35) // 36)
        ws.row_dimensions[r].height = max(22, 14 * description_lines)
        fill = FILL_WHITE if idx % 2 == 0 else FILL_LIGHT_GRAY

        # ITEM
        c_item = ws.cell(row=r, column=2, value=idx + 1)
        c_item.font = BODY_FONT
        c_item.alignment = CENTER_NOWRAP
        c_item.fill = copy(fill)
        c_item.border = THIN_BORDER

        # PROVEEDOR
        c_prov = ws.cell(row=r, column=3)
        _escribir_texto_literal(c_prov, g.get("proveedor", ""))
        c_prov.font = BODY_FONT
        c_prov.alignment = LEFT
        c_prov.fill = copy(fill)
        c_prov.border = THIN_BORDER

        # FECHA (guardar como texto dd/mm/yyyy para no pelear con locales)
        fecha_val = g.get("fecha", "")
        if isinstance(fecha_val, (datetime, date)):
            fecha_val = fecha_val.strftime("%d/%m/%Y")
        c_fecha = ws.cell(row=r, column=4)
        _escribir_texto_literal(c_fecha, fecha_val)
        c_fecha.font = BODY_FONT
        c_fecha.alignment = CENTER_NOWRAP
        c_fecha.fill = copy(fill)
        c_fecha.border = THIN_BORDER

        # FACTURA
        c_fact = ws.cell(row=r, column=5)
        _escribir_texto_literal(c_fact, g.get("factura", ""))
        c_fact.font = BODY_FONT
        c_fact.alignment = CENTER_NOWRAP
        c_fact.fill = copy(fill)
        c_fact.border = THIN_BORDER
        c_fact.number_format = "@"

        # TOTAL Bs
        monto = g.get("monto_bs", None)
        try:
            monto_f = float(str(monto).replace(".", "").replace(",", ".")) if isinstance(monto, str) and "," in str(monto) else (float(monto) if monto not in (None, "") else None)
        except (ValueError, TypeError):
            monto_f = None
        c_bs = ws.cell(row=r, column=6, value=monto_f)
        c_bs.font = BODY_FONT
        c_bs.alignment = RIGHT
        c_bs.fill = copy(fill)
        c_bs.border = THIN_BORDER
        c_bs.number_format = fmt_bs

        # TOTAL $ = Bs / tasa (fórmula viva; si tasa=0 queda 0 para no dividir por cero)
        tasa_coord = "$F$8"
        c_usd = ws.cell(row=r, column=7)
        c_usd.value = f"=IF({tasa_coord}=0,0,F{r}/{tasa_coord})"
        c_usd.font = BODY_FONT
        c_usd.alignment = RIGHT
        c_usd.fill = copy(fill)
        c_usd.border = THIN_BORDER
        c_usd.number_format = fmt_usd

        # DESCRIPCIÓN
        c_desc = ws.cell(row=r, column=8)
        _escribir_texto_literal(c_desc, g.get("descripcion", ""))
        c_desc.font = BODY_FONT
        c_desc.alignment = LEFT
        c_desc.fill = copy(fill)
        c_desc.border = THIN_BORDER

    # --- Totales ---
    total_row = end_row + 1
    ws.row_dimensions[total_row].height = 22
    ws.merge_cells(f"B{total_row}:E{total_row}")
    ws[f"B{total_row}"] = tr(lang, "excel_total_expenses")
    ws[f"B{total_row}"].font = TOTAL_FONT
    ws[f"B{total_row}"].alignment = Alignment(horizontal="right", vertical="center")
    _style_range(ws, total_row, 2, 5, fill=FILL_NAVY)
    for c in range(2, 6):
        ws.cell(row=total_row, column=c).border = THIN_BORDER
        ws.cell(row=total_row, column=c).font = TOTAL_FONT

    ws[f"F{total_row}"] = f"=SUM(F{start_row}:F{end_row})"
    ws[f"F{total_row}"].font = TOTAL_FONT
    ws[f"F{total_row}"].fill = copy(FILL_NAVY)
    ws[f"F{total_row}"].alignment = RIGHT
    ws[f"F{total_row}"].border = THIN_BORDER
    ws[f"F{total_row}"].number_format = fmt_bs

    ws[f"G{total_row}"] = f"=SUM(G{start_row}:G{end_row})"
    ws[f"G{total_row}"].font = TOTAL_FONT
    ws[f"G{total_row}"].fill = copy(FILL_NAVY)
    ws[f"G{total_row}"].alignment = RIGHT
    ws[f"G{total_row}"].border = THIN_BORDER
    ws[f"G{total_row}"].number_format = fmt_usd
    ws[f"H{total_row}"].fill = copy(FILL_NAVY)
    ws[f"H{total_row}"].border = THIN_BORDER

    sr = total_row + 1
    ws.row_dimensions[sr].height = 22
    ws.merge_cells(f"B{sr}:E{sr}")
    ws[f"B{sr}"] = tr(lang, "excel_balance")
    ws[f"B{sr}"].font = Font(name="Calibri", size=10, bold=True, color=NAVY)
    ws[f"B{sr}"].alignment = Alignment(horizontal="right", vertical="center")
    _style_range(ws, sr, 2, 5, fill=FILL_LIGHT_BLUE)
    for c in range(2, 6):
        ws.cell(row=sr, column=c).border = THIN_BORDER
    ws.merge_cells(f"F{sr}:F{sr}")
    ws[f"F{sr}"] = "—"
    ws[f"F{sr}"].alignment = CENTER
    ws[f"F{sr}"].border = THIN_BORDER
    ws[f"G{sr}"] = f"=$C$8-G{total_row}"
    ws[f"G{sr}"].font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    ws[f"G{sr}"].alignment = RIGHT
    ws[f"G{sr}"].border = THIN_BORDER
    ws[f"G{sr}"].number_format = fmt_usd
    ws[f"H{sr}"].border = THIN_BORDER

    # --- Firmas ---
    fr = sr + 3
    ws.merge_cells(f"B{fr}:C{fr}")
    ws.merge_cells(f"D{fr}:E{fr}")
    ws.merge_cells(f"F{fr}:H{fr}")
    signature_headers = (
        (f"B{fr}", tr(lang, "excel_prepared_by")),
        (f"D{fr}", tr(lang, "excel_reviewed_by")),
        (f"F{fr}", tr(lang, "excel_approved_by")),
    )
    for coord, txt in signature_headers:
        ws[coord] = txt
        ws[coord].font = Font(name="Calibri", size=9, bold=True, color=WHITE)
        ws[coord].alignment = CENTER
        ws[coord].fill = copy(FILL_BLUE)
        ws[coord].border = THIN_BORDER
    ws.row_dimensions[fr].height = 20

    for _ in range(3):
        fr += 1
        ws.row_dimensions[fr].height = 14
    ws.merge_cells(f"B{fr}:C{fr}")
    ws.merge_cells(f"D{fr}:E{fr}")
    ws.merge_cells(f"F{fr}:H{fr}")
    for col in ("B", "D", "F"):
        ws[f"{col}{fr}"].border = Border(bottom=Side(style="medium", color=NAVY))
        ws[f"{col}{fr}"].alignment = CENTER
    fr += 1
    ws.merge_cells(f"B{fr}:C{fr}")
    ws.merge_cells(f"D{fr}:E{fr}")
    ws.merge_cells(f"F{fr}:H{fr}")
    signature_date = tr(lang, "excel_signature_date")
    ws[f"B{fr}"] = signature_date
    ws[f"D{fr}"] = signature_date
    ws[f"F{fr}"] = signature_date
    for col in ("B", "D", "F"):
        ws[f"{col}{fr}"].font = SMALL_FONT
        ws[f"{col}{fr}"].alignment = CENTER

    # --- Ajustes de hoja ---
    ws.freeze_panes = "B11"
    ws.auto_filter.ref = f"B10:H{end_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.print_title_rows = "10:10"

    wb.properties.creator = tr(lang, "app_title")
    wb.properties.title = f"{tr(lang, 'excel_sheet_title')} — {responsable} — {fecha_emision}"
    wb.properties.company = nombre
    return wb
